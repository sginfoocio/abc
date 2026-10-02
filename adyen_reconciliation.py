"""Offline readers and transformations for Adyen settlement reports."""

import csv
import hashlib
import io
import re
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation


class ReportError(ValueError):
    """A file cannot be interpreted safely."""


@dataclass(frozen=True)
class SourceRow:
    row: int
    values: dict[str, str]


def _read_csv(content: bytes) -> list[tuple[int, list]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252")
    dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=",;\t")
    reader = csv.reader(io.StringIO(text), dialect)
    return [(reader.line_num, values) for values in reader]


def _read_xlsx(content: bytes) -> list[tuple[int, list]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        return [(number, list(values)) for number, values in enumerate(workbook.active.values, start=1)]
    finally:
        workbook.close()


def _read_xls(content: bytes) -> list[tuple[int, list]]:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=content)
    try:
        sheet = workbook.sheet_by_index(0)
        return [(row_index + 1, [
            xlrd.xldate_as_datetime(cell.value, workbook.datemode)
            if cell.ctype == xlrd.XL_CELL_DATE else cell.value
            for cell in sheet.row(row_index)
        ]) for row_index in range(sheet.nrows)]
    finally:
        workbook.release_resources()


def _cell_text(value) -> str:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return "" if value is None else str(value).strip()


def read_table(content: bytes, filename: str) -> list[SourceRow]:
    suffix = filename.rsplit(".", 1)[-1].lower()
    reader = {"csv": _read_csv, "xlsx": _read_xlsx, "xls": _read_xls}.get(suffix)
    if reader is None:
        raise ReportError("Formato no admitido: use XLS, XLSX o CSV.")
    try:
        records = reader(content)
    except Exception as exc:
        raise ReportError("Archivo ilegible, vacio o incompatible con su extension.") from exc

    nonempty = [(number, values) for number, values in records if any(value not in (None, "") for value in values)]
    if not nonempty:
        raise ReportError("El archivo no contiene una cabecera.")
    _, header_values = nonempty[0]
    headers = [str(value or "").strip() for value in header_values]
    if any(not header for header in headers) or len(set(headers)) != len(headers):
        raise ReportError("Cabeceras vacias o duplicadas.")
    rows = []
    for number, values in nonempty[1:]:
        if len(values) != len(headers):
            raise ReportError(f"Fila {number}: numero de columnas distinto de la cabecera.")
        normalized = {header: _cell_text(value) for header, value in zip(headers, values)}
        rows.append(SourceRow(number, normalized))
    return rows


def money(value: str, *, blank_zero: bool = False) -> Decimal:
    text = value.strip()
    if not text and blank_zero:
        return Decimal("0")
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    try:
        amount = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("Importe invalido (sin separadores de miles).") from exc
    if not amount.is_finite():
        raise ValueError("Importe no finito.")
    return amount


@dataclass(frozen=True)
class RowError:
    source: str
    row: int
    reason: str


@dataclass(frozen=True)
class Movement:
    row: int
    account: str
    batch: str
    currency: str
    date: str
    kind: str
    net: Decimal
    gross: Decimal | None
    original_currency: str
    psp: str
    order: str
    expenses: tuple[tuple[str, Decimal], ...]


@dataclass(frozen=True)
class Document:
    row: int
    number: str
    partner: str
    date: str
    amount: Decimal
    currency: str
    psp: str
    order: str


ADYEN_COLUMNS = {
    "account": ("Merchant Account", "Account"),
    "batch": ("Batch Number", "Batch"),
    "date": ("Booking Date", "Creation Date", "Date"),
    "kind": ("Type",),
    "currency": ("Currency", "Net Currency"),
    "original_currency": ("Gross", "Gross Currency"),
    "net_credit": ("Net Credit (NC)",),
    "net_debit": ("Net Debit (NC)",),
    "gross_credit": ("Gross Credit (GC)",),
    "gross_debit": ("Gross Debit (GC)",),
    "psp": ("Psp Reference", "PSP Reference"),
    "order": ("Merchant Reference", "Order Reference"),
}
ODOO_COLUMNS = {
    "number": ("Número", "Numero", "Number"),
    "partner": ("Nombre del socio a mostrar en la factura.", "Nombre del socio a mostrar en la factura", "Invoice Partner Display Name"),
    "date": ("Fecha de Factura/Recibo", "Invoice/Bill Date"),
    "amount": ("Total en divisa firmado", "Total Signed in Currency"),
    "currency": ("Moneda", "Currency"),
    "psp": ("PSP Reference", "Psp Reference", "Referencia PSP"),
    "order": ("Merchant Reference", "Order Reference", "Referencia de pedido", "Origen", "Documento origen", "Source Document"),
}
EXPENSE_COLUMNS = ("Commission (NC)", "Markup (NC)", "Scheme Fees (NC)", "Interchange (NC)")
MOVEMENT_TYPES = {"Settled", "Refunded", "Fee", "InvoiceDeduction", "MerchantPayout"}
CONFIRMED_DOCUMENT = "Documento confirmado"


def _columns(rows: list[SourceRow], aliases: dict, required: tuple[str, ...]) -> dict[str, str]:
    if not rows:
        raise ReportError("El archivo no contiene movimientos o documentos.")
    headers = rows[0].values
    columns = {
        field: next((name for name in names if name in headers), "")
        for field, names in aliases.items()
    }
    missing = [aliases[field][0] for field in required if not columns[field]]
    if missing:
        raise ReportError("Faltan columnas: " + "; ".join(missing))
    return columns


def _value(row: SourceRow, columns: dict, field: str) -> str:
    return row.values.get(columns[field], "")


def _date(value: str) -> str:
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        for pattern in ("%Y-%m-%d %H:%M:%S %Z", "%d/%m/%Y", "%d/%m/%Y %H:%M:%S"):
            try:
                return datetime.strptime(value, pattern).date().isoformat()
            except ValueError:
                pass
    raise ValueError("Fecha invalida; use ISO o dia/mes/ano.")


def _currency(value: str) -> str:
    currency = value.upper()
    if not re.fullmatch(r"[A-Z]{3}", currency):
        raise ValueError("Moneda invalida; se requiere codigo ISO de tres letras.")
    return currency


def _gross_amount(values: dict[str, str]) -> tuple[Decimal | None, str]:
    if values["gross_credit"] or values["gross_debit"]:
        gross = money(values["gross_credit"], blank_zero=True) - money(values["gross_debit"], blank_zero=True)
        return gross, _currency(values["original_currency"])
    if values["kind"] in {"Settled", "Refunded"}:
        raise ValueError("Falta importe bruto o moneda original para el cobro/devolucion.")
    return None, ""


def validate_movements(rows: list[SourceRow]) -> tuple[list[Movement], list[RowError]]:
    columns = _columns(rows, ADYEN_COLUMNS, (
        "account", "batch", "date", "kind", "currency", "net_credit", "net_debit",
    ))
    movements, errors = [], []
    for row in rows:
        try:
            values = {field: _value(row, columns, field) for field in columns}
            if not values["account"] or not values["batch"]:
                raise ValueError("Cuenta o lote vacio.")
            if values["kind"] not in MOVEMENT_TYPES:
                raise ValueError("Tipo de movimiento no admitido; no se descarta silenciosamente.")
            if not values["net_credit"] and not values["net_debit"]:
                raise ValueError("Credito y debito netos vacios.")
            net = money(values["net_credit"], blank_zero=True) - money(values["net_debit"], blank_zero=True)
            gross, original_currency = _gross_amount(values)
            expenses = tuple(
                (name, money(row.values[name])) for name in EXPENSE_COLUMNS
                if row.values.get(name)
            )
            movements.append(Movement(
                row.row, values["account"], values["batch"], _currency(values["currency"]),
                _date(values["date"]), values["kind"], net, gross, original_currency,
                values["psp"], values["order"], expenses,
            ))
        except ValueError as exc:
            errors.append(RowError("Adyen", row.row, str(exc)))
    return movements, errors


def validate_documents(rows: list[SourceRow]) -> tuple[list[Document], list[RowError]]:
    columns = _columns(rows, ODOO_COLUMNS, ("number", "partner", "date", "amount", "currency"))
    documents, errors = [], []
    for row in rows:
        try:
            values = {field: _value(row, columns, field) for field in columns}
            if not values["number"]:
                raise ValueError("Numero de documento vacio.")
            documents.append(Document(
                row.row, values["number"], values["partner"], _date(values["date"]),
                money(values["amount"]), _currency(values["currency"]),
                values["psp"], values["order"],
            ))
        except ValueError as exc:
            errors.append(RowError("Odoo", row.row, str(exc)))
    return documents, errors


def decimal_text(amount: Decimal) -> str:
    text = format(amount, "f")
    if "." not in text:
        return text + ".00"
    return text + "0" if len(text.rsplit(".", 1)[1]) == 1 else text


def _reference_match(movement: Movement, document: Document) -> bool:
    return bool(
        (movement.psp and movement.psp == document.psp)
        or (movement.order and movement.order == document.order)
    )


def _match_status(matches: list[int], reverse: dict[int, list[int]] | list[list[int]], confirmed: bool) -> str:
    if confirmed:
        return "confirmada"
    if not matches:
        return "sin_coincidencia"
    return "candidata" if len(matches) == 1 and len(reverse[matches[0]]) == 1 else "ambigua"


def _review_record(source: str, row: int, status: str, number: str, related: list[int], candidates: list[str], amount, currency: str) -> dict:
    return {
        "Origen": source, "Fila": row, "Estado": status, CONFIRMED_DOCUMENT: number,
        "Filas relacionadas": ", ".join(str(other) for other in related),
        "Documentos candidatos": ", ".join(candidates),
        "Bruto firmado": decimal_text(amount) if amount is not None else "", "Moneda original": currency,
    }


def match_documents(movements: list[Movement], documents: list[Document]) -> tuple[list[dict], list[dict]]:
    by_amount = defaultdict(list)
    for index, document in enumerate(documents):
        by_amount[(document.amount, document.currency)].append(index)
    candidates = [
        by_amount.get((movement.gross, movement.original_currency), [])
        if movement.kind in {"Settled", "Refunded"} else []
        for movement in movements
    ]
    reverse = defaultdict(list)
    reference_edges = []
    for index, matches in enumerate(candidates):
        reference_edges.append([other for other in matches if _reference_match(movements[index], documents[other])])
        for other in matches:
            reverse[other].append(index)
    reference_usage = Counter(other for matches in reference_edges for other in matches)
    confirmed = {
        index: matches[0] for index, matches in enumerate(reference_edges)
        if len(matches) == 1 and reference_usage[matches[0]] == 1
    }
    movement_review = []
    for index, movement in enumerate(movements):
        matches = candidates[index]
        status = _match_status(matches, reverse, index in confirmed)
        number = documents[confirmed[index]].number if index in confirmed else ""
        movement_review.append(_review_record(
            "Adyen", movement.row, status, number, [documents[other].row for other in matches],
            [documents[other].number for other in matches], movement.gross, movement.original_currency,
        ))
    document_review = []
    confirmed_documents = set(confirmed.values())
    for index, document in enumerate(documents):
        matches = reverse[index]
        status = _match_status(matches, candidates, index in confirmed_documents)
        document_review.append(_review_record(
            "Odoo", document.row, status, document.number if status == "confirmada" else "",
            [movements[other].row for other in matches], [document.number], document.amount, document.currency,
        ))
    return movement_review, document_review


def concept(movement: Movement, confirmed_number: str = "") -> str:
    parts = [f"Cuenta {movement.account}", f"Lote {movement.batch}", movement.kind, f"fila {movement.row}"]
    if movement.psp:
        parts.append(f"PSP {movement.psp}")
    if movement.order:
        parts.append(f"Pedido {movement.order}")
    if movement.original_currency and movement.original_currency != movement.currency:
        parts.append(f"Original {decimal_text(movement.gross)} {movement.original_currency}")
    if movement.expenses:
        parts.append("Gastos informados (incluidos en neto; no sumar): " + ", ".join(
            f"{name}={decimal_text(amount)} {movement.currency}" for name, amount in movement.expenses
        ))
    if confirmed_number:
        parts.append(f"Documento {confirmed_number}")
    return " | ".join(parts)


def _csv_bytes(headers: list[str], rows: list[list]) -> bytes:
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
    writer.writerow(headers)
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


def build_exports(movements: list[Movement], review: list[dict]) -> dict[str, bytes]:
    groups = defaultdict(list)
    for movement, match in zip(movements, review, strict=True):
        groups[(movement.account, movement.batch, movement.currency)].append([
            movement.date, concept(movement, match[CONFIRMED_DOCUMENT]), decimal_text(movement.net),
        ])
    exports = {}
    for key, rows in groups.items():
        safe_key = "_".join(re.sub(r"[^A-Za-z0-9_-]", "_", value)[:60] for value in key)
        digest = hashlib.sha256(repr(key).encode()).hexdigest()[:10]
        exports[f"adyen_{safe_key}_{digest}.csv"] = _csv_bytes(["Fecha", "Concepto", "Importe"], rows)
    return exports


@dataclass
class ReconciliationResult:
    movements: list[Movement]
    documents: list[Document]
    errors: list[RowError]
    movement_review: list[dict]
    document_review: list[dict]
    exports: dict[str, bytes]

    def review_csv(self) -> bytes:
        records = self.movement_review + self.document_review
        if not records:
            return b""
        headers = list(records[0])
        return _csv_bytes(headers, [[record[header] for header in headers] for record in records])

    def zip_bytes(self) -> bytes:
        if self.errors or not self.exports:
            raise ReportError("No se pueden descargar resultados con errores o sin movimientos.")
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            for filename, content in self.exports.items():
                archive.writestr(filename, content)
            archive.writestr("revision.csv", self.review_csv())
        return output.getvalue()

    def totals(self) -> list[dict]:
        groups = defaultdict(list)
        for movement in self.movements:
            groups[(movement.account, movement.batch, movement.currency)].append(movement)
        return [{
            "Cuenta": account, "Lote": batch, "Moneda": currency, "Movimientos": len(items),
            "Neto": decimal_text(sum((item.net for item in items), Decimal(0))),
            "MerchantPayout": decimal_text(sum((item.net for item in items if item.kind == "MerchantPayout"), Decimal(0))),
        } for (account, batch, currency), items in groups.items()]


def process_reports(adyen: bytes, filename: str, odoo: bytes | None = None, odoo_filename: str = "") -> ReconciliationResult:
    movements, errors = validate_movements(read_table(adyen, filename))
    documents = []
    if odoo is not None:
        documents, document_errors = validate_documents(read_table(odoo, odoo_filename))
        errors.extend(document_errors)
    movement_review, document_review = match_documents(movements, documents)
    exports = build_exports(movements, movement_review) if not errors else {}
    return ReconciliationResult(movements, documents, errors, movement_review, document_review, exports)